"""Добавляет артикул продавца на PDF-стикеры WB. aiogram 2.25.1."""
import asyncio
import logging
import re
import tempfile
from pathlib import Path

import fitz
from aiogram import types
from aiogram.dispatcher import FSMContext
from aiogram.dispatcher.filters.state import State, StatesGroup
from openpyxl import load_workbook

from loader import bot, dp

LOG = logging.getLogger(__name__)
if not LOG.handlers:
    handler = logging.FileHandler('sticker_errors.log', encoding='utf-8')
    handler.setFormatter(logging.Formatter('%(asctime)s %(levelname)s %(message)s'))
    LOG.addHandler(handler)
LOG.setLevel(logging.ERROR)
MAX_BYTES = 20 * 1024 * 1024


class StickerFiles(StatesGroup):
    waiting = State()


def normalize_header(value):
    return re.sub(r'\s+', ' ', str(value or '').strip()).casefold()


def parse_number(value, row_num):
    if isinstance(value, bool) or value is None:
        raise ValueError(f'Excel, строка {row_num}: номер стикера пустой или неверный.')
    if isinstance(value, (int, float)):
        if not float(value).is_integer():
            raise ValueError(f'Excel, строка {row_num}: номер стикера не целый.')
        value = str(int(value))
    key = re.sub(r'\s+', '', str(value))
    if not re.fullmatch(r'[0-9]{11}', key):
        raise ValueError(f'Excel, строка {row_num}: номер стикера {value!r} должен содержать 11 цифр.')
    return key


def parse_excel(path):
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError('XLSX повреждён или имеет неподдерживаемый формат.') from exc
    try:
        for ws in wb:
            ws.reset_dimensions()  # Иногда WB записывает неверную область листа.
            rows = ws.iter_rows(values_only=True)
            for row_num, row in enumerate(rows, 1):
                if row_num > 30:
                    break
                headers = {normalize_header(v): i for i, v in enumerate(row) if v is not None}
                article_col = headers.get('артикул продавца')
                number_col = headers.get('стикер')
                if number_col is None:
                    number_col = headers.get('стикер (номер)')
                if number_col is None:
                    number_col = headers.get('стикер wb')
                if number_col is None:
                    number_col = headers.get('стикер/номер')
                if article_col is None or number_col is None:
                    continue
                mapping = {}
                for data_row_num, values in enumerate(rows, row_num + 1):
                    article = values[article_col] if article_col < len(values) else None
                    number = values[number_col] if number_col < len(values) else None
                    if article is None and number is None:
                        continue
                    if article is None or str(article).strip() == '':
                        raise ValueError(f'Excel, строка {data_row_num}: не заполнен артикул продавца.')
                    key = parse_number(number, data_row_num)
                    value = str(article).strip()
                    if not value.isascii() or not value.isprintable():
                        raise ValueError(f'Excel, строка {data_row_num}: артикул {value!r} содержит неподдерживаемые символы.')
                    if key in mapping:
                        raise ValueError(f'Excel, строка {data_row_num}: повторяется стикер {key}.')
                    mapping[key] = value
                if not mapping:
                    raise ValueError(f'На листе «{ws.title}» нет заполненных товаров.')
                return mapping

        # Старый экспорт «Лист подбора»: заголовки в строке 5, номер в D.
        for ws in wb:
            ws.reset_dimensions()
            rows = ws.iter_rows(values_only=True)
            for row_num, row in enumerate(rows, 1):
                if row_num > 30:
                    break
                headers = {normalize_header(v): i for i, v in enumerate(row) if v is not None}
                if 'артикул продавца' not in headers:
                    continue
                num_col = headers.get('стикер')
                if num_col is None:
                    num_col = headers.get('стикер (номер)')
                if num_col is None:
                    num_col = headers.get('стикер wb')
                if num_col is None and 'штрихкод' in headers:
                    # «Штрихкод» не является номером стикера: игнорируем.
                    pass
                if num_col is None and 'лист подбора' in ws.title.casefold():
                    num_col = 3
                if num_col is None:
                    continue
                mapping = {}
                for data_row_num, values in enumerate(rows, row_num + 1):
                    ac = headers['артикул продавца']
                    article = values[ac] if ac < len(values) else None
                    number = values[num_col] if num_col < len(values) else None
                    if article is None and number is None:
                        continue
                    if article is None or not str(article).strip():
                        raise ValueError(f'Excel, строка {data_row_num}: артикул продавца не заполнен.')
                    key = parse_number(number, data_row_num)
                    value = str(article).strip()
                    if not value.isascii() or not value.isprintable():
                        raise ValueError(f'Excel, строка {data_row_num}: неподдерживаемые символы в артикуле.')
                    if key in mapping:
                        raise ValueError(f'Excel, строка {data_row_num}: повтор номера {key}.')
                    mapping[key] = value
                if mapping:
                    return mapping
        raise ValueError('Не найдены колонки «Стикер» и «Артикул продавца». Нужен экспорт WB с этими полями.')
    finally:
        wb.close()


def make_stickers(pdf_path, excel_path, output_path):
    mapping = parse_excel(excel_path)
    try:
        pdf = fitz.open(pdf_path)
    except Exception as exc:
        raise ValueError('PDF повреждён или имеет неверный формат.') from exc
    try:
        if pdf.needs_pass:
            raise ValueError('PDF защищён паролем.')
        if not pdf.page_count:
            raise ValueError('PDF не содержит страниц.')
        found = []
        for index, page in enumerate(pdf, 1):
            if page.rotation != 0 or not (110 <= page.rect.width <= 116 and 82 <= page.rect.height <= 88):
                raise ValueError(f'Страница {index}: неподдерживаемый размер или поворот стикера.')
            lines = [x.strip() for x in page.get_text('text', sort=False).splitlines()]
            ids = [a + b for a, b in zip(lines, lines[1:])
                   if re.fullmatch(r'[0-9]{7}', a) and re.fullmatch(r'[0-9]{4}', b)]
            if len(ids) != 1:
                raise ValueError(f'Страница {index}: невозможно однозначно определить номер стикера.')
            key = ids[0]
            if key not in mapping:
                raise ValueError(f'Страница {index}: номер {key} отсутствует в Excel.')
            if key in found:
                raise ValueError(f'Страница {index}: номер {key} повторяется в PDF.')
            found.append(key)
        extra = set(mapping) - set(found)
        if extra:
            raise ValueError(f'В Excel {len(mapping)} товаров, в PDF {len(found)} стикеров. '
                             f'Нет стикеров для {len(extra)} номеров (например, {min(extra)}).')

        # Координаты проверены для исходного стикера 40x30 мм.
        for page, key in zip(pdf, found):
            article = mapping[key]
            unit_width = fitz.get_text_length(article, fontname='hebo', fontsize=1)
            size = min(7.8, 34.0 / max(unit_width, 0.01))
            if size < 2.7:
                raise ValueError(f'Артикул {article!r} слишком длинный для свободной зоны стикера.')
            page.insert_text((8.3, 59.5), article, fontname='hebo', fontsize=size,
                             rotate=90, color=(0, 0, 0), overlay=True)
        pdf.save(output_path, garbage=3, deflate=True)
        return len(found)
    finally:
        pdf.close()


@dp.message_handler(commands=['cancel'], state='*')
async def cancel_stickers(message: types.Message, state: FSMContext):
    await state.finish()
    await message.answer('Загрузка отменена. Отправьте PDF и XLSX заново.')


@dp.message_handler(content_types=types.ContentType.DOCUMENT, state='*')
async def sticker_document(message: types.Message, state: FSMContext):
    ext = Path(message.document.file_name or '').suffix.lower()
    if ext not in ('.pdf', '.xlsx'):
        return await message.answer('Отправьте PDF со стикерами и XLSX с артикулами.')
    if message.document.file_size and message.document.file_size > MAX_BYTES:
        return await message.answer('Файл больше 20 МБ.')
    data = await state.get_data()
    if ext in data:
        return await message.answer(f'{ext.upper()} уже получен. Для замены используйте /cancel.')
    await state.update_data(**{ext: message.document.file_id,
                               f'{ext}_name': message.document.file_name or 'stickers.pdf'})
    await StickerFiles.waiting.set()
    data = await state.get_data()
    if '.pdf' not in data or '.xlsx' not in data:
        return await message.answer(f'Файл получен. Теперь отправьте {"XLSX" if ext == ".pdf" else "PDF"}.')

    try:
        await message.answer('Проверяю номера и формирую PDF...')
        with tempfile.TemporaryDirectory() as temp:
            src_pdf = str(Path(temp) / 'input.pdf')
            src_excel = str(Path(temp) / 'input.xlsx')
            dst_pdf = str(Path(temp) / 'result.pdf')
            await bot.download_file_by_id(data['.pdf'], src_pdf)
            await bot.download_file_by_id(data['.xlsx'], src_excel)
            count = await asyncio.to_thread(make_stickers, src_pdf, src_excel, dst_pdf)
            with open(dst_pdf, 'rb') as file:
                await message.answer_document(
                    types.InputFile(file, filename=f"{Path(data['.pdf_name']).stem}_added.pdf"),
                    caption=f'Готово: {count} из {count} стикеров.'
                )
    except ValueError as exc:
        LOG.exception('user_id=%s: ошибка данных', message.from_user.id)
        await message.answer(f'Не удалось создать PDF. Причина: {exc}\nОтправьте исправленные файлы заново.')
    except Exception:
        LOG.exception('user_id=%s: системная ошибка', message.from_user.id)
        await message.answer('Техническая ошибка обработки или отправки PDF. Подробности в sticker_errors.log.')
    finally:
        await state.finish()
